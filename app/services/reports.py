from __future__ import annotations

import asyncio
import os
import uuid
from datetime import UTC, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.models import AccountJobResult, Job, TelegramAccount
from app.repositories import aggregate_job, job_results
from app.services.control import aware


def masked_phone(value):
    return value if len(value) < 6 else value[:3] + "*" * (len(value) - 5) + value[-2:]


def append_text_safe(sheet, values):
    row = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
    for col, value in enumerate(values, 1):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
        cell = sheet.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = "s"


def render_workbook(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            append_text_safe(sheet, row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24364B")
            cell.alignment = Alignment(wrap_text=True)
        for column in sheet.columns:
            width = min(55, max(12, max(len(str(cell.value or "")) for cell in column[:200]) + 2))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    temporary = path.with_name(path.name + ".tmp")
    try:
        with temporary.open("xb") as stream:
            os.chmod(temporary, 0o600)
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


class ReportService:
    def __init__(self, settings, db):
        self.settings, self.db = settings, db
        self._slots = asyncio.Semaphore(2)

    async def _write(self, prefix, sheets):
        self.settings.reports_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
        path = self.settings.reports_dir / f"{prefix}_{uuid.uuid4().hex}.xlsx"
        async with self._slots:
            await asyncio.to_thread(render_workbook, path, sheets)
        return path

    def _date(self, value):
        return (
            aware(value).astimezone(self.settings.tz).strftime("%d.%m.%Y %H:%M:%S") if value else ""
        )

    async def build_job_report(self, job_id, *, mask_phones=False):
        async with self.db.session() as session:
            job = await session.get(Job, job_id)
            if job is None:
                raise ValueError("Задание не найдено")
            aggregates = await aggregate_job(session, job_id)
            results = await job_results(session, job_id)
        summary = [
            [
                "Канал",
                "Ссылка",
                "Аккаунтов успешно",
                "Аккаунтов запланировано",
                "Обработано публикаций (макс.)",
                "Запросов регистрации просмотра",
            ]
        ]
        summary += [
            [
                r["title"],
                r["url"],
                r["accounts_completed"],
                r["accounts_total"],
                r["available_posts"],
                r["viewed_posts"],
            ]
            for r in aggregates
        ]
        details = [
            [
                "Задание",
                "Канал",
                "Ссылка",
                "№ аккаунта",
                "Телефон",
                "Результат",
                "Этап",
                "Обработано публикаций",
                "Последний ID",
                "Прокси ID",
                "Начало",
                "Завершение",
                "Тип причины",
                "Причина",
            ]
        ]
        for r in results:
            details.append(
                [
                    job.id,
                    r.channel_title,
                    r.channel_url,
                    r.account_number,
                    masked_phone(r.account_phone) if mask_phones else r.account_phone,
                    r.status.value,
                    r.state,
                    r.viewed_posts,
                    r.last_message_id,
                    r.proxy_id or "",
                    self._date(r.started_at),
                    self._date(r.finished_at),
                    r.error_kind or "",
                    r.error or "",
                ]
            )
        problems = [details[0]] + [
            row
            for row, result in zip(details[1:], results, strict=True)
            if result.status.value != "completed"
        ]
        info = [
            ["Параметр", "Значение"],
            ["Статус", job.status.value],
            ["Создал", str(job.created_by)],
            ["Создано", self._date(job.created_at)],
            ["Начало", self._date(job.started_at)],
            ["Завершение", self._date(job.finished_at)],
            ["Хранить до", self._date(job.retention_until)],
            ["Результатов запланировано", job.expected_results],
            ["Результатов записано", len(results)],
            [
                "Метрика",
                "Успешные запросы API не гарантируют прирост публичного счётчика Telegram.",
            ],
            ["Телефоны", "Скрыты для этой выгрузки" if mask_phones else "Полные номера"],
        ]
        if (job.summary or {}).get("legacy_planning_unknown"):
            info.append(
                [
                    "Старое задание",
                    "Импортировано из 1.2: полный исходный план неизвестен; количества отражают только сохранившиеся результаты.",
                ]
            )
        path = await self._write(
            f"report_{job.id}",
            [
                ("Итог по каналам", summary),
                ("Подробно по аккаунтам", details),
                ("Проблемные результаты", problems),
                ("О задании", info),
            ],
        )
        if not mask_phones:
            async with self.db.session() as session:
                row = await session.get(Job, job.id)
                if row:
                    row.report_path = str(path)
        return path

    async def build_problem_accounts_report(self, *, mask_phones=False):
        async with self.db.session() as session:
            rows = list(
                (
                    await session.scalars(
                        select(TelegramAccount)
                        .where(
                            TelegramAccount.archived_at.is_(None),
                            TelegramAccount.last_error.is_not(None),
                        )
                        .order_by(TelegramAccount.display_number)
                    )
                ).all()
            )
        values = [["№", "Телефон", "Статус", "Последняя проверка", "Причина"]]
        values += [
            [
                r.display_number,
                masked_phone(r.phone) if mask_phones else r.phone,
                r.status.value,
                self._date(r.health_checked_at),
                r.last_error or "",
            ]
            for r in rows
        ]
        return await self._write("problem_accounts", [("Проблемные аккаунты", values)])

    async def daily(self, day, channel_ids=None):
        start = datetime.combine(day, time.min, tzinfo=self.settings.tz).astimezone(UTC)
        end = datetime.combine(
            day + timedelta(days=1), time.min, tzinfo=self.settings.tz
        ).astimezone(UTC)
        async with self.db.session() as session:
            results = list(
                (
                    await session.scalars(
                        select(AccountJobResult)
                        .join(Job)
                        .where(Job.created_at >= start, Job.created_at < end)
                        .where(
                            AccountJobResult.channel_id.in_(channel_ids)
                            if channel_ids is not None
                            else True
                        )
                    )
                ).all()
            )
        totals = {}
        for row in results:
            data = totals.setdefault(
                row.channel_url,
                {
                    "title": row.channel_title,
                    "jobs": set(),
                    "accounts": set(),
                    "views": 0,
                    "failures": 0,
                },
            )
            data["jobs"].add(row.job_id)
            data["accounts"].add(row.account_id)
            data["views"] += row.viewed_posts
            data["failures"] += row.status.value in ("failed", "cancelled")
        rows = [
            [
                "Канал",
                "Ссылка",
                "Запусков",
                "Уникальных аккаунтов",
                "Запросов регистрации просмотра",
                "Ошибок/отмен",
            ]
        ]
        rows += [
            [r["title"], url, len(r["jobs"]), len(r["accounts"]), r["views"], r["failures"]]
            for url, r in totals.items()
        ]
        return await self._write(f"daily_{day.isoformat()}", [("Сводка дня", rows)])

    async def subscription_report(self, result):
        rows = [["№ аккаунта", "Результат", "Причина"]]
        rows += [[number, "Успешно", ""] for number in result.succeeded]
        rows += [[number, "Ошибка", reason] for number, reason in result.failed]
        return await self._write("subscriptions", [("Результат", rows)])
